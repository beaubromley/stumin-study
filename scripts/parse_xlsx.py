"""Parse the StuMin MASTER Scope and Sequence workbook into data/curriculum.json.

One entry per school year (worksheet), each holding an ordered list of blocks:
volume headers, unit headers, and sessions. Session hyperlinks on the title cell
are the Google Drive lesson PDFs.
"""
import datetime
import json
import os
import re
import sys

import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\beaub\Downloads\StuMin MASTER Scope and Sequence.xlsx"
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "curriculum.json")

NO_MEETING = {"no cgs", "no cg", "skipped"}


def norm(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() or None


def drive_id(url):
    m = re.search(r"/d/([^/]+)", url or "")
    return m.group(1) if m else None


def parse_sheet(ws):
    blocks = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=6):
        a, b, c, d, e, f = (norm(x.value) for x in row)
        if not any((a, b, c, d, e, f)):
            continue
        title_cell = row[3]
        link = title_cell.hyperlink.target if title_cell.hyperlink else None

        if a and a.lower().startswith("volume"):
            blocks.append({"kind": "volume", "title": a})
        elif a and a.lower().startswith("key passage"):
            text = c or b or d
            if text:
                blocks.append({"kind": "key_passage", "text": text})
        elif (a and a.lower() == "unit") or (a is None and c is None and d and b):
            blocks.append({"kind": "unit", "number": b, "title": d})
        elif b and b.lower() in NO_MEETING or (a and a.lower() in NO_MEETING and not d):
            blocks.append({"kind": "break", "date": a, "label": b or a})
        elif d:
            blocks.append(
                {
                    "kind": "session",
                    "date": a,
                    "unit": b,
                    "session": c,
                    "title": d,
                    "main_point": e,
                    "passage": f,
                    "pdf_url": link,
                    "drive_id": drive_id(link),
                }
            )
    return blocks


def main():
    wb = openpyxl.load_workbook(SRC)
    years = []
    for ws in wb.worksheets:
        blocks = parse_sheet(ws)
        sessions = [x for x in blocks if x["kind"] == "session"]
        years.append(
            {
                "year": ws.title,
                "blocks": blocks,
                "session_count": len(sessions),
                "linked_count": sum(1 for s in sessions if s["pdf_url"]),
            }
        )
    years.sort(key=lambda y: y["year"], reverse=True)
    with open(OUT, "w", encoding="utf8") as fh:
        json.dump({"years": years}, fh, ensure_ascii=False, indent=1)
    for y in years:
        print(f"{y['year']}: {y['session_count']} sessions, {y['linked_count']} with PDFs")


if __name__ == "__main__":
    main()
